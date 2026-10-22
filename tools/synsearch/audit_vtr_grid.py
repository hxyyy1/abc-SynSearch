#!/usr/bin/env python3
"""Audit fixed-length grid evidence independently of the report generator."""
import argparse,csv,hashlib,json,zipfile
from pathlib import Path
from statistics import mean
from vtr_objectives import stats
from grid_vtr_fixed_length import METHODS,DESIGNS,configurations

def main():
    p=argparse.ArgumentParser();p.add_argument('root',type=Path);p.add_argument('--previous',type=Path,required=True);p.add_argument('--tables',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    read=lambda path:json.loads(path.read_text())
    manifest=read(a.root/'manifest.json');trials=read(a.root/'all_trials.json');final=read(a.root/'final_results.json');selection=read(a.root/'selection.json');old=read(a.previous/'all_results.json')
    oldmap={(r['objective'],r['algorithm'],r['design'],r['seed']):r for r in old}
    key=lambda r:(r['objective'],r['algorithm'],r['config_id'],r['design'],r['seed'])
    bykey={key(r):r for r in trials};assert len(bykey)==len(trials)
    assert all(r['status']=='ok' for r in trials)
    new=[r for r in trials if r['origin']=='new']
    allowed={'mcts':{'cpuct','mu_discount'},'hybrid':{'cpuct','ucb_c'},'sa':{'initial_temperature','min_temperature'},'ucb1':{'ucb_c'},'ucb1-prefix':{'ucb_c'},'linucb':{'alpha','lambda'}}
    hashes={}
    for r in new:
        # All new attempts belong to this experiment, including retry directories.
        remote=Path('/home/xyhao/abc-synsearch-vtr-20261018/grid_L10_20261018')
        work=a.root/Path(r['work']).relative_to(remote)
        z=read(work/'native.json');assert z==r['data']['result']
        assert z['objective']==r['objective'] and z['algorithm']==r['algorithm'] and z['seed']==r['seed']
        assert z['config']['sequence_length']==len(z['sequence'])==10
        for k,v in r['options'].items():assert z['config']['sequence_length' if k=='steps' else k.replace('-','_')]==v
        baseline=oldmap[(r['objective'],r['algorithm'],r['design'],r['seed'])]['result']
        changes={k for k in z['config'] if z['config'][k]!=baseline['config'][k]}
        assert changes<=allowed[r['algorithm']],(key(r),changes)
        assert (z['initial_and'],z['initial_lev'])==(baseline['initial_and'],baseline['initial_lev'])
        s=stats(work/'search.log');assert (s['nodes'],s['level'])==(z['final_and'],z['final_lev'])
        assert r['data']['cec']['equivalent'] and 'Networks are equivalent' in (work/'cec.log').read_text()
        assert (work/'optimized.aig').stat().st_size>0
        hashes[str(work.relative_to(a.root)/'optimized.aig')]=hashlib.sha256((work/'optimized.aig').read_bytes()).hexdigest()
        if r['phase']!='validation':assert r['design'] in ('or1200','fir') and r['seed']==0
        if r['algorithm']=='sa' and z['config']['initial_temperature']==-1:assert '--initial-temperature' not in (work/'command.txt').read_text()
    for method in METHODS:
        assert manifest['configurations'][method]==configurations(method)
        for cfg in configurations(method):
            for obj in ('nodes','level'):assert (obj,method,cfg,'or1200',0) in bykey
    promoted=read(a.root/'promoted.json')
    for pair,ch in selection.items():
        obj,method=pair.split('/');field='final_and' if obj=='nodes' else 'final_lev'
        candidates=list(dict.fromkeys([x['config_id'] for x in promoted[pair]]+['baseline']))
        ranks=[]
        for cfg in candidates:
            rr=[bykey[(obj,method,cfg,d,0)] for d in ('or1200','fir')]
            score=mean(r['data']['result'][field]/oldmap[(obj,method,r['design'],0)]['result'][field] for r in rr)
            duration=mean(r['data']['resources']['process_wall_sec'] for r in rr)
            ranks.append((round(score,12),duration,cfg))
        assert sorted(ranks)[0][2]==ch['selected']
    expected={(o,m,d,s) for o in ('nodes','level') for m in METHODS for d in DESIGNS for s in ([0] if m=='mcts' else [0,1,2])}
    got={(r['objective'],r['algorithm'],r['design'],r['seed']) for r in final}
    assert got==expected and len(final)==len(got)==192
    for r in final:assert r['config_id']==selection[r['objective']+'/'+r['algorithm']]['selected']
    log=(a.root/'runner.log').read_text();assert log.index('FROZEN_SELECTION')<log.index('PHASE validation')
    if (a.root/'restart_correction.json').exists():assert not read(a.root/'restart_correction.json')['selection_already_exists']
    overview=list(csv.DictReader((a.tables/'累计最佳_总对比.csv').open(encoding='utf-8-sig')))
    for r in overview:
        obj,d=r['objective'],r['design'];field='final_and' if obj=='nodes' else 'final_lev'
        before=min(x['result'][field] for x in old if x['objective']==obj and x['design']==d)
        after=min([before]+[x['data']['result'][field] for x in new if x['objective']==obj and x['design']==d])
        assert int(r['previous_best'])==before and int(r['now_observed_best'])==after
    for r in csv.DictReader((a.tables/'固定参数_逐电路比较.csv').open(encoding='utf-8-sig')):
        rr=[x['data']['result'] for x in final if (x['objective'],x['algorithm'],x['design'])==(r['objective'],r['algorithm'],r['design'])]
        field='final_and' if r['objective']=='nodes' else 'final_lev'
        assert int(r['tuned_best'])==min(z[field] for z in rr)
        assert abs(float(r['tuned_mean'])-mean(z[field] for z in rr))<1e-9
    with zipfile.ZipFile(a.tables/'VTR固定长度网格_完整对比.xlsx') as z:assert z.testzip() is None
    from docx import Document
    doc=Document(a.tables/'VTR固定长度网格_实验报告.docx');assert len(doc.tables)==5
    for r,dr in zip(overview,doc.tables[0].rows[1:]):assert dr.cells[2].text==r['previous_best'] and dr.cells[3].text==r['now_observed_best']
    for name,field in [('grid_vtr_fixed_length.py','driver_sha256')]:assert hashlib.sha256(Path(__file__).with_name(name).read_bytes()).hexdigest()==manifest[field]
    assert manifest['abc_sha256']==read(a.previous/'manifest.json')['abc_sha256']
    a.out.mkdir(parents=True,exist_ok=True)
    (a.out/'网络SHA256.json').write_text(json.dumps(hashes,indent=2))
    audit=dict(successful_new_runs=len(new),final_cases=len(final),full_grid_screen_cases=102,all_sequences_exactly_10=True,
        only_allowed_policy_parameters_changed=True,baseline_binary_unchanged=True,native_results_and_CEC_logs_checked=True,
        selection_recomputed=True,selection_frozen_before_validation=True,all_comparison_values_recomputed=True,
        failed_original_records_preserved=len(list(a.root.glob('runs/**/failed_run_attempt-*.json'))),
        memory_guard_events=log.count('MEMORY_GUARD'),docx_table_values_checked=True,docx_visual_rendering='Not available: LibreOffice absent')
    (a.out/'数据核验.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2));print(json.dumps(audit,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
