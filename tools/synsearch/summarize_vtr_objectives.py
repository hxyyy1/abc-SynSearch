#!/usr/bin/env python3
"""Validate and summarize vtr_objectives.py results into CSV, XLSX and Word.
Run with a Python environment containing python-docx and XlsxWriter.
"""
import argparse
import csv
import json
from pathlib import Path
from statistics import mean, stdev

METHODS=('mcts','hybrid','sa','ucb1','ucb1-prefix','linucb')
LABELS={'mcts':'MCTS','hybrid':'HybridSyn','sa':'SA','ucb1':'全局 UCB1','ucb1-prefix':'位置 UCB1','linucb':'LinUCB'}
DESIGNS=('bfly','dscg','fir','ode','or1200','syn2')

def export_csv(path,rows):
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)


def main():
    ap=argparse.ArgumentParser();ap.add_argument('results',type=Path);ap.add_argument('--out',type=Path,required=True)
    args=ap.parse_args();root=args.results;out=args.out;out.mkdir(parents=True,exist_ok=True)
    manifest=json.loads((root/'manifest.json').read_text())
    records=json.loads((root/'all_results.json').read_text())
    bases={n:json.loads((root/'baselines'/n/'baseline.json').read_text()) for n in DESIGNS}
    expected={(o,a,d,s) for o in ('nodes','level') for a in METHODS for d in DESIGNS for s in manifest['seeds']}
    observed=[(r['objective'],r['algorithm'],r['design'],r['seed']) for r in records]
    assert len(observed)==len(set(observed)) and set(observed)==expected
    failures=[r for r in records if r['status']!='ok']
    if failures:raise RuntimeError(f'Incomplete/failed runs: {len(failures)}; inspect raw statuses')
    assert all(r['cec']['equivalent'] for r in records)
    assert all(b['cec']['equivalent'] for b in bases.values())
    summaries=[];raw=[];lookup={}
    for obj in ('nodes','level'):
        primary='final_and' if obj=='nodes' else 'final_lev'
        secondary='final_lev' if obj=='nodes' else 'final_and'
        for name in DESIGNS:
            base=bases[name]
            for method in METHODS:
                group=sorted((r for r in records if (r['objective'],r['design'],r['algorithm'])==(obj,name,method)),key=lambda r:r['seed'])
                best=min(group,key=lambda r:(r['result'][primary],r['result'][secondary],r['seed']))
                result=best['result'];values=[r['result'][primary] for r in group]
                initial=base['initial']['nodes' if obj=='nodes' else 'level']
                baseline=base['resyn2']['nodes' if obj=='nodes' else 'level']
                row=dict(objective=obj,design=name,algorithm=method,n_seeds=len(group),initial_target=initial,
                         resyn2_target=baseline,best_target=min(values),mean_target=mean(values),sd_target=stdev(values),
                         best_seed=best['seed'],best_nodes=result['final_and'],best_level=result['final_lev'],
                         reduction_vs_initial_pct=100*(initial-min(values))/initial,
                         reduction_vs_resyn2_pct=100*(baseline-min(values))/baseline,
                         mean_search_sec=mean(r['result']['total_runtime_sec'] for r in group),
                         mean_process_sec=mean(r['resources']['process_wall_sec'] for r in group),
                         mean_evaluations=mean(r['result']['evaluations'] for r in group),
                         mean_peak_rss_mib=mean(r['resources']['peak_rss_kib']/1024 for r in group),
                         cec_passed=f'{len(group)}/{len(group)}',best_recipe=result['recipe'])
                summaries.append(row);lookup[(obj,name,method)]=row
                for r in group:
                    z=r['result'];raw.append(dict(objective=obj,design=name,algorithm=method,seed=r['seed'],
                        initial_and=z['initial_and'],initial_lev=z['initial_lev'],final_and=z['final_and'],final_lev=z['final_lev'],
                        target_value=z[primary],search_sec=z['total_runtime_sec'],process_sec=r['resources']['process_wall_sec'],
                        evaluations=z['evaluations'],cache_hits=z['cache_hits'],peak_rss_mib=r['resources']['peak_rss_kib']/1024,
                        cec='PASS',recipe=z['recipe']))
    export_csv(out/'逐电路算法汇总.csv',summaries);export_csv(out/'全部216次实验.csv',raw)
    matrices={};aggregate=[]
    for obj in ('nodes','level'):
        metric='nodes' if obj=='nodes' else 'level'
        rows=[]
        for name in DESIGNS:
            row={'电路':name,'初始值':bases[name]['initial'][metric],'resyn2':bases[name]['resyn2'][metric]}
            for method in METHODS:row[LABELS[method]]=lookup[(obj,name,method)]['best_target']
            smallest=min(lookup[(obj,name,m)]['best_target'] for m in METHODS)
            row['搜索最佳算法']=' / '.join(LABELS[m] for m in METHODS if lookup[(obj,name,m)]['best_target']==smallest)
            rows.append(row)
        matrices[obj]=rows;export_csv(out/('节点数对比.csv' if obj=='nodes' else 'Level对比.csv'),rows)
        for method in METHODS:
            data=[lookup[(obj,n,method)] for n in DESIGNS]
            aggregate.append(dict(objective=obj,algorithm=method,
                mean_best_reduction_pct=mean(r['reduction_vs_initial_pct'] for r in data),
                mean_seed_reduction_pct=mean(100*(r['initial_target']-r['mean_target'])/r['initial_target'] for r in data),
                wins_vs_resyn2=sum(r['best_target']<r['resyn2_target'] for r in data),
                ties_vs_resyn2=sum(r['best_target']==r['resyn2_target'] for r in data),
                worse_vs_resyn2=sum(r['best_target']>r['resyn2_target'] for r in data),
                mean_process_sec=mean(r['mean_process_sec'] for r in data),
                mean_peak_rss_mib=mean(r['mean_peak_rss_mib'] for r in data)))
    export_csv(out/'算法总体汇总.csv',aggregate)
    # Editable workbook; primary tables show best of three and detailed sheets
    # expose mean/std, the corresponding secondary metric, and all seed results.
    import xlsxwriter
    book=xlsxwriter.Workbook(out/'VTR六电路_双目标实验对比.xlsx')
    title=book.add_format({'bold':True,'font_size':15,'font_color':'#17365D'})
    header=book.add_format({'bold':True,'bg_color':'#17365D','font_color':'white','text_wrap':True,'border':1})
    cell=book.add_format({'border':1,'valign':'top'});decimal=book.add_format({'border':1,'num_format':'0.000'})
    win=book.add_format({'bg_color':'#DDF2DF','bold':True})
    note=book.add_format({'text_wrap':True,'valign':'top'})
    def sheet(name,rows):
        ws=book.add_worksheet(name);keys=list(rows[0]);ws.freeze_panes(1,3);ws.autofilter(0,0,len(rows),len(keys)-1)
        for j,k in enumerate(keys):ws.write(0,j,k,header);ws.set_column(j,j,19 if j<18 else 48)
        for i,row in enumerate(rows,1):
            for j,k in enumerate(keys):
                value=row[k];ws.write(i,j,value,decimal if isinstance(value,float) else cell)
        ws.set_row(0,32);ws.set_landscape();ws.fit_to_pages(1,0);ws.repeat_rows(0)
        return ws
    cover=book.add_worksheet('说明与实验配置');cover.set_column(0,0,27);cover.set_column(1,1,100)
    cover.write(0,0,'VTR 六电路双目标实验',title)
    notes=[('实验规模',f'6 电路 × 6 算法 × 2 目标 × {len(manifest["seeds"])} 种子 = {len(records)} 次；另有 6 次 resyn2。'),
        ('主表口径','每个算法三个种子的最小目标值，不是全局最优证明；对应另一指标和种子见详细汇总。'),
        ('种子与序列长度',f'{manifest["seeds"]}；长度 {manifest["parameters"]["sequence_length"]}'),
        ('算法预算',json.dumps(manifest['parameters'],ensure_ascii=False)),
        ('动作集合',manifest['actions']),
        ('硬件','AMD Ryzen Threadripper PRO 7985WX，3 个独立物理核并发，单个 ABC 进程绑定一个核。'),
        ('资源限制','用户 MemoryHigh=12 GiB、MemoryMax=16 GiB；六路高预算预运行因换页中止，已排除于正式统计。'),
        ('时间','process_sec 为完整 ABC 搜索进程耗时，不含单独 CEC；search_sec 为原生计时。算法预算不等时。'),
        ('正确性',f'{len(records)} 次搜索和 6 次 resyn2 的 CEC 均通过；保持寄存器边界的组合逻辑等价检查。'),
        ('代码变更','新增 --objective nodes|level；默认 legacy 不变；切换 MCTS/Hybrid 奖励及 LinUCB 尾段和回溯判据。节点暴涨保护仍保留。'),
        ('局限','只有六个 VTR 电路与三个种子；结果是有限预算启发式搜索；AIG 指标不是映射面积或实际时延。'),
        ('二进制 SHA256',manifest['abc_sha256']),('源码 SHA256',manifest['source_sha256']),('远端目录','/home/xyhao/abc-synsearch-vtr-20261018/results')]
    for i,(k,v) in enumerate(notes,2):cover.write(i,0,k,header);cover.write(i,1,v,note);cover.set_row(i,48)
    for obj,name in [('nodes','节点最小化_最佳值'),('level','Level最小化_最佳值')]:
        ws=sheet(name,matrices[obj]);ws.set_column(9,9,65)
        ws.conditional_format(1,3,6,8,{'type':'formula','criteria':'=D2=MIN($D2:$I2)','format':win})
        rows=[{'电路':n,**{LABELS[m]:f"{lookup[(obj,n,m)]['mean_target']:.2f} ± {lookup[(obj,n,m)]['sd_target']:.2f}" for m in METHODS}} for n in DESIGNS]
        sheet('节点均值与标准差' if obj=='nodes' else 'Level均值与标准差',rows)
    sheet('算法总体汇总',aggregate);sheet('逐电路72组汇总',summaries);sheet('全部216次实验',raw)
    base_rows=[dict(design=n,**{f'initial_{k}':v for k,v in bases[n]['initial'].items()},
                   **{f'resyn2_{k}':v for k,v in bases[n]['resyn2'].items()},
                   process_sec=bases[n]['resources']['process_wall_sec'],cec='PASS') for n in DESIGNS]
    sheet('原始及resyn2基线',base_rows);book.close()
    # Word report, no unmeasured results or diagrams.
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(29.7);sec.page_height=Cm(21)
    sec.top_margin=sec.bottom_margin=Cm(1.6);sec.left_margin=sec.right_margin=Cm(1.5)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(10);s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('17365D')
    doc.styles['Heading 2'].font.size=Pt(13)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    def p(t):doc.add_paragraph(t)
    def heading(t):doc.add_heading(t,1)
    def tab(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if i==0:trpr.append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.space_before=Pt(3)
                    for r in pp.runs:r.font.size=Pt(8);r.bold=i==0
        return t
    heading('ABC-SynSearch：VTR 六电路双目标实验报告')
    p('实验目标：分别最小化 AIG 节点数与逻辑层数（level）。本报告可作为课程论文的 VTR 实验补充，EPFL 实验仍未开展。')
    for k,v in notes[:10]:p(k+'：'+v)
    doc.add_page_break();heading('1  最小化节点数：六种算法对比')
    keys=list(matrices['nodes'][0]);tab(keys,[list(r.values()) for r in matrices['nodes']])
    p('表 1：每格为该算法三个种子中获得的最少节点数；绿色标记仅用于 Excel，表示六种搜索算法中的并列最优。resyn2 是固定十步脚本，耗时远低于多轮搜索。')
    node_winners={n:min(lookup[('nodes',n,m)]['best_target'] for m in METHODS) for n in DESIGNS}
    for n in DESIGNS:
        a=bases[n]['initial']['nodes'];b=bases[n]['resyn2']['nodes'];v=node_winners[n]
        p(f'{n}：搜索最佳 {v:,} 个节点，较初始 {a:,} 下降 {(a-v)/a*100:.2f}%，较 resyn2 的 {b:,} 下降 {(b-v)/b*100:.2f}%。')
    doc.add_page_break();heading('2  最小化 level：六种算法对比')
    keys=list(matrices['level'][0]);tab(keys,[list(r.values()) for r in matrices['level']])
    p('表 2：每格为三个种子中获得的最小 level。以另一指标打破并列时仅决定选用哪条同目标值轨迹，不改变该表主指标。')
    level_winners={n:min(lookup[('level',n,m)]['best_target'] for m in METHODS) for n in DESIGNS}
    for n in DESIGNS:
        a=bases[n]['initial']['level'];b=bases[n]['resyn2']['level'];v=level_winners[n]
        p(f'{n}：搜索最佳 level={v}，较初始 {a} 下降 {(a-v)/a*100:.2f}%；resyn2 为 {b}，差值为 {b-v} 层。')
    doc.add_page_break();heading('3  随机种子稳定性与计算成本')
    for obj,label in [('nodes','节点数'),('level','level')]:
        doc.add_heading(label+'：均值 ± 样本标准差',2)
        tab(['电路']+[LABELS[m] for m in METHODS],[[n]+[f"{lookup[(obj,n,m)]['mean_target']:.2f} ± {lookup[(obj,n,m)]['sd_target']:.2f}" for m in METHODS] for n in DESIGNS])
    p('每组 n=3。MCTS 当前策略为确定性，修改种子可得到相同轨迹，其跨种子标准差不能反映所有构建或环境变化。主表是 best-of-three，不能将其当作平均性能。')
    doc.add_page_break();heading('4  算法总体统计与解释边界')
    tab(['目标','算法','最佳下降率均值/%','跨种子平均下降率/%','胜/平/负 resyn2','平均进程耗时/s','平均峰值RSS/MiB'],
        [[r['objective'],LABELS[r['algorithm']],f"{r['mean_best_reduction_pct']:.2f}",f"{r['mean_seed_reduction_pct']:.2f}",
          f"{r['wins_vs_resyn2']}/{r['ties_vs_resyn2']}/{r['worse_vs_resyn2']}",f"{r['mean_process_sec']:.2f}",f"{r['mean_peak_rss_mib']:.1f}"] for r in aggregate])
    p('下降率先相对各电路初始值归一化，再对六个电路取等权平均。耗时与内存为该算法在六电路、三个种子的运行均值；不包含事后的独立 CEC。并发任务绑定不同物理核，但仍可能共享缓存和内存带宽，因此不能将此表当作严格独占机器的性能测试。')
    p('由于 MCTS 每步展开、SA 完整序列提案、UCB 回合和 LinUCB 前瞻的成本不同，表中只比较明确预算下的工程结果，不声称等时间或等评估量排名。若需要效率结论，应进一步扫描预算并画出质量-时间关系。')
    doc.add_page_break();heading('5  实验实现、正确性与复现')
    p('原实现 MCTS/Hybrid 固定 5:2 混合奖励，LinUCB 尾段和回溯偏向节点。为满足本次纯目标实验，新增 --objective nodes|level：切换 MCTS 奖励、Hybrid 分支排序、SA/UCB 权重、LinUCB 即时奖励、尾段/深度最佳比较及回溯差距。默认 legacy 保留原始策略；节点暴涨限制保留为约束。此处结果属于目标适配版本，应区别于此前课程报告中的默认算法语义。')
    p('远端构建后 11 项集成测试通过，包括六算法默认轨迹回归，以及以一步穷举为独立参照的双目标选择测试。全部 216 次搜索均通过初始/输出统计核对、寄存器数量检查和 CEC；6 个 resyn2 基线也通过 CEC。CEC 比较的是保持寄存器边界的组合逻辑功能，不是重新编码或重定时后的通用时序等价证明。')
    p('远端实验目录：/home/xyhao/abc-synsearch-vtr-20261018。原始数据在 results/，包括各次输入/输出 AIG、原生命令与 JSON、搜索日志、资源日志和独立 CEC 日志。manifest.json 记录二进制、源码、驱动与输入 SHA-256，以及硬件和参数。')
    p('复现命令：python3 tools/synsearch/vtr_objectives.py --out /path/to/new/results --workers 3。汇总命令：python3 tools/synsearch/summarize_vtr_objectives.py /path/to/results --out /path/to/tables。')
    p('二进制 SHA-256：'+manifest['abc_sha256']);p('核心源码 SHA-256：'+manifest['source_sha256'])
    p('所有最小值均为指定动作集、长度、预算和种子下观测到的最好结果，不代表理论最优。AIG 节点数与 level 仍是结构指标，不等于技术映射后的面积和时序。')
    doc.save(out/'VTR六电路_实验结果报告.docx')
    summary={'runs':len(records),'cec_passed':len(records),'baseline_cec_passed':6,'matrices':matrices,'aggregate':aggregate}
    (out/'汇总.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
