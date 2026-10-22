#!/usr/bin/env python3
"""Export locked-config validation and separately labeled best-observed results."""
import argparse,csv,json,hashlib
from collections import defaultdict
from pathlib import Path
from statistics import mean,stdev
from tune_vtr import METHODS,DESIGNS,TRAIN,HOLDOUT
from summarize_vtr_objectives import LABELS,export_csv


def main():
    ap=argparse.ArgumentParser();ap.add_argument('tuning',type=Path);ap.add_argument('--previous',type=Path,required=True);ap.add_argument('--out',type=Path,required=True)
    a=ap.parse_args();root=a.tuning;out=a.out;out.mkdir(parents=True,exist_ok=True)
    manifest=json.loads((root/'manifest.json').read_text());selected=json.loads((root/'selection.json').read_text())
    final=json.loads((root/'final_results.json').read_text());trials=json.loads((root/'all_trials.json').read_text())
    old=json.loads((a.previous/'all_results.json').read_text());oldmap={(r['objective'],r['algorithm'],r['design'],r['seed']):r for r in old}
    expected={(o,m,d,s) for o in ('nodes','level') for m in METHODS for d in DESIGNS for s in ([0] if m=='mcts' else [0,1,2])}
    keys=[(r['objective'],r['algorithm'],r['design'],r['seed']) for r in final]
    assert len(keys)==len(set(keys)) and set(keys)==expected
    failures=[r for r in final if r['status']!='ok']
    if failures:raise RuntimeError(f'{len(failures)} locked-config runs failed; report those before producing complete comparison')
    for r in trials:
        if r['status']=='ok':assert r['data']['cec']['equivalent']
    groups=defaultdict(list)
    for r in final:groups[(r['objective'],r['algorithm'],r['design'])].append(r)
    detail=[];incumbents=[];raw=[]
    for (obj,m,d),rs in sorted(groups.items()):
        field='final_and' if obj=='nodes' else 'final_lev';other='final_lev' if obj=='nodes' else 'final_and'
        rs.sort(key=lambda r:r['seed']);controls=[oldmap[(obj,m,d,r['seed'])] for r in rs]
        vals=[r['data']['result'][field] for r in rs];before=[r['result'][field] for r in controls]
        best=min(rs,key=lambda r:(r['data']['result'][field],r['data']['result'][other],r['seed']))
        q=best['data']['result'];cfg=selected[f'{obj}/{m}'];n=len(rs)
        fresh=[r for r in rs if r['origin']=='new']
        row=dict(objective=obj,algorithm=m,design=d,split='train' if d in TRAIN else 'holdout',
            selected_config=cfg['selected'],sequence_length=cfg['options']['steps'],n=n,
            baseline_best=min(before),tuned_best=min(vals),best_gain=min(before)-min(vals),
            best_gain_pct=100*(min(before)-min(vals))/min(before),baseline_mean=mean(before),tuned_mean=mean(vals),
            mean_gain_pct=100*(mean(before)-mean(vals))/mean(before),tuned_sd=stdev(vals) if n>1 else None,
            improved_seeds=sum(x<y for x,y in zip(vals,before)),tied_seeds=sum(x==y for x,y in zip(vals,before)),
            worse_seeds=sum(x>y for x,y in zip(vals,before)),best_seed=best['seed'],best_nodes=q['final_and'],best_level=q['final_lev'],
            baseline_mean_sec=mean(r['resources']['process_wall_sec'] for r in controls),
            tuned_mean_sec=mean(r['data']['resources']['process_wall_sec'] for r in rs),
            tuned_mean_peak_mib=mean(r['data']['resources']['peak_rss_kib']/1024 for r in rs),
            new_runs=len(fresh),cec_passed=f'{n}/{n}',recipe=q['recipe'])
        detail.append(row)
        choices=[]
        for seed in (0,1,2):
            r=oldmap[(obj,m,d,seed)]
            choices.append(dict(config_id='baseline',seed=seed,result=r['result'],origin='previous',path=str(a.previous/'runs'/obj/m/d/f'seed-{seed}'/'optimized.aig')))
        for r in trials:
            if (r['objective'],r['algorithm'],r['design'])==(obj,m,d) and r['status']=='ok' and r['origin']=='new':
                choices.append(dict(config_id=r['config_id'],seed=r['seed'],result=r['data']['result'],origin='new',path=str(Path(r['work'])/'optimized.aig')))
        winner=min(choices,key=lambda r:(r['result'][field],r['result'][other],r['config_id'],r['seed']))
        prior=min(oldmap[(obj,m,d,s)]['result'][field] for s in (0,1,2))
        z=winner['result']
        incumbents.append(dict(objective=obj,algorithm=m,design=d,baseline_best=prior,
            observed_best=z[field],gain=prior-z[field],gain_pct=100*(prior-z[field])/prior,
            config_id=winner['config_id'],seed=winner['seed'],sequence_length=z['config']['sequence_length'],
            nodes=z['final_and'],level=z['final_lev'],trials_available=len(choices),origin=winner['origin'],
            network_path=winner['path'],recipe=z['recipe']))
    for r in trials:
        row=dict(phase=r['phase'],objective=r['objective'],algorithm=r['algorithm'],config_id=r['config_id'],
                 design=r['design'],seed=r['seed'],origin=r['origin'],status=r['status'])
        if r['status']=='ok':
            z=r['data']['result'];row.update(nodes=z['final_and'],level=z['final_lev'],length=z['config']['sequence_length'],
                 evaluations=z['evaluations'],process_sec=r['data']['resources']['process_wall_sec'],
                 peak_rss_mib=r['data']['resources']['peak_rss_kib']/1024,cec='PASS',recipe=z['recipe'])
        else:row.update(nodes=None,level=None,length=r['options']['steps'],evaluations=None,process_sec=None,peak_rss_mib=None,cec='NOT PASSED',recipe='')
        raw.append(row)
    raw.sort(key=lambda r:(r['objective'],r['algorithm'],r['config_id'],r['design'],r['seed']))
    aggregate=[]
    for obj in ('nodes','level'):
        for m in METHODS:
            for split in ('all','holdout','train'):
                rr=[r for r in detail if r['objective']==obj and r['algorithm']==m and (split=='all' or r['split']==split)]
                aggregate.append(dict(objective=obj,algorithm=m,split=split,config=selected[f'{obj}/{m}']['selected'],
                     circuits=len(rr),mean_best_gain_pct=mean(r['best_gain_pct'] for r in rr),
                     mean_mean_gain_pct=mean(r['mean_gain_pct'] for r in rr),
                     improved_circuits=sum(r['best_gain']>0 for r in rr),tied_circuits=sum(r['best_gain']==0 for r in rr),
                     worse_circuits=sum(r['best_gain']<0 for r in rr),
                     baseline_mean_sec=mean(r['baseline_mean_sec'] for r in rr),tuned_mean_sec=mean(r['tuned_mean_sec'] for r in rr)))
    cfgrows=[]
    for obj in ('nodes','level'):
        for m in METHODS:
            z=selected[f'{obj}/{m}'];cfgrows.append(dict(objective=obj,algorithm=m,selected_config=z['selected'],
                train_ratio=z['ranking'][0]['mean_ratio'],options=json.dumps(z['options'],ensure_ascii=False,sort_keys=True)))
    lookup={(r['objective'],r['algorithm'],r['design']):r for r in incumbents}
    overview=[];matrices={}
    for obj in ('nodes','level'):
        matrices[obj]=[]
        for d in DESIGNS:
            rrs=[lookup[(obj,m,d)] for m in METHODS];before=min(r['baseline_best'] for r in rrs);after=min(r['observed_best'] for r in rrs)
            winners=[r for r in rrs if r['observed_best']==after]
            overview.append(dict(objective=obj,design=d,previous_best=before,now_observed_best=after,gain=before-after,
                 gain_pct=100*(before-after)/before,
                 winners=' / '.join(LABELS[r['algorithm']] for r in winners),
                 winner_configs=' / '.join(r['algorithm']+':'+r['config_id'] for r in winners)))
            matrices[obj].append({'电路':d,**{LABELS[m]:lookup[(obj,m,d)]['observed_best'] for m in METHODS}})
    export_csv(out/'固定参数_逐电路比较.csv',detail);export_csv(out/'固定参数_总体与留出集.csv',aggregate)
    export_csv(out/'累计最佳_逐算法电路.csv',incumbents);export_csv(out/'累计最佳_总对比.csv',overview)
    export_csv(out/'全部配置试验.csv',raw);export_csv(out/'选定参数.csv',cfgrows)
    for obj,label in [('nodes','节点'),('level','Level')]:export_csv(out/f'{label}_累计最佳对比.csv',matrices[obj])
    monitor=[json.loads(s) for s in (root/'memory_monitor.jsonl').read_text().splitlines()]
    completed=json.loads((root/'completed.json').read_text())
    memory=dict(max_active=max(s['active'] for s in monitor),max_user_memory_gib=max(s['memory_bytes'] for s in monitor)/2**30,
                min_swap_mib=min(s['swap_bytes'] for s in monitor)/2**20,max_swap_mib=max(s['swap_bytes'] for s in monitor)/2**20,
                high_gib=manifest['memory_high_bytes']/2**30,max_gib=manifest['memory_max_bytes']/2**30)
    notes=[('方法','逐轮淘汰配置竞赛：or1200 seed0 初筛 5 配置（含旧配置）；前两名在 fir seed0 复核，并保留旧配置对照；按两训练电路平均目标比值选定一组参数。'),
           ('验证边界','参数在验证前锁定。bfly、dscg、ode、syn2 为未参与选参的电路；随机方法使用 seed0/1/2，训练电路另用 seed1/2 检查。MCTS 无随机策略，仅运行 seed0。'),
           ('两种表的区别','固定参数表用于比较选定配置的表现，允许变差；累计最佳表保留此前与本轮所有有效结果中的最好值，按构造不会变差，不能当作泛化性能。'),
           ('预算','候选同时包含策略参数、更大预算以及长度20/较少单步预算；不是等长度或等算力竞赛。完整选项见选定参数与原始日志。'),
           ('并行执行',f"最多 {memory['max_active']} 个独立 ABC 进程实际并发，固定物理核；历史峰值估计进行内存准入，用户内存观测最大 {memory['max_user_memory_gib']:.2f} GiB。"),
           ('正确性','沿用上一轮同一 ABC 二进制。每个有效新结果均做 CEC 与寄存器数量核对；组合逻辑等价检查保持寄存器边界。'),
           ('耗时解释','旧结果为上轮三路并发测量，新结果采用动态并发；可以描述实际成本，不能作为严格隔离环境的加速比。'),
           ('数据规模',json.dumps(completed,ensure_ascii=False)),
           ('复现位置','/home/xyhao/abc-synsearch-vtr-20261018/tuning_20261018'),
           ('二进制SHA256',manifest['abc_sha256'])]
    import xlsxwriter
    wb=xlsxwriter.Workbook(out/'VTR调参_完整对比.xlsx')
    head=wb.add_format({'bold':True,'bg_color':'#17365D','font_color':'white','text_wrap':True,'border':1})
    cell=wb.add_format({'border':1,'valign':'top'});num=wb.add_format({'border':1,'num_format':'0.000'})
    wrap=wb.add_format({'text_wrap':True,'valign':'top'})
    def sheet(name,rows):
        ws=wb.add_worksheet(name);keys=list(rows[0]);ws.freeze_panes(1,3);ws.autofilter(0,0,len(rows),len(keys)-1)
        for j,k in enumerate(keys):ws.write(0,j,k,head);ws.set_column(j,j,20)
        for i,r in enumerate(rows,1):
            for j,k in enumerate(keys):ws.write(i,j,r[k],num if isinstance(r[k],float) else cell)
        ws.set_row(0,35);return ws
    ws=wb.add_worksheet('说明');ws.set_column(0,0,20);ws.set_column(1,1,120)
    for i,(k,v) in enumerate(notes):ws.write(i,0,k,head);ws.write(i,1,v,wrap);ws.set_row(i,60)
    sheet('累计最佳_总对比',overview);sheet('节点_累计最佳',matrices['nodes']);sheet('Level_累计最佳',matrices['level'])
    ws=sheet('选定参数',cfgrows);ws.set_column(4,4,130)
    sheet('固定参数_留出集',[r for r in aggregate if r['split']=='holdout'])
    sheet('固定参数_全体',aggregate);sheet('固定参数_72组',detail)
    ws=sheet('累计最佳_72组',incumbents);ws.set_column(14,15,60)
    sheet('全部配置试验',raw);wb.close()
    from docx import Document
    from docx.shared import Cm,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();s=doc.sections[0];s.page_width=Cm(29.7);s.page_height=Cm(21)
    s.left_margin=s.right_margin=Cm(1.5);s.top_margin=s.bottom_margin=Cm(1.6)
    for name in ('Normal','Heading 1','Heading 2'):
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(10);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(13)
    def p(t):doc.add_paragraph(t)
    def h(t):doc.add_heading(t,1)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.space_after=Pt(3)
                    for r in pp.runs:r.font.size=Pt(8);r.bold=i==0
    h('ABC-SynSearch：VTR 参数调优与验证报告')
    for k,v in notes[:7]:p(k+'：'+v)
    p('完成统计：'+json.dumps(completed,ensure_ascii=False))
    doc.add_page_break();h('1  是否找到更好的结果：累计最佳对比')
    table(['目标','电路','上轮最好','本轮累计最好','减少量','改善/%','获得该值的算法'],
        [[r['objective'],r['design'],r['previous_best'],r['now_observed_best'],r['gain'],f"{r['gain_pct']:.2f}",r['winners']] for r in overview])
    p('本表跨算法、参数配置与种子选取已通过验证的最好网络，回答“是否找到更好的电路”。它不是固定调优配置的独立测试结果；训练电路经历更多配置尝试。所有最优值仅指当前搜索范围内已观测的值。')
    doc.add_page_break();h('2  六种算法各自的累计最好结果')
    for obj,label in [('nodes','节点数'),('level','逻辑层数')]:
        doc.add_heading(label,2);table(list(matrices[obj][0]),[list(r.values()) for r in matrices[obj]])
    p('每格允许保留上轮结果，数值不变不意味着没有运行调参。每个单元格的配置编号、种子、网络路径及另一指标保存在 Excel 的“累计最佳_72组”中。')
    doc.add_page_break();h('3  锁定参数在四个留出电路上的表现')
    held=[r for r in aggregate if r['split']=='holdout']
    table(['目标','算法','选定配置','最好值平均改善/%','均值平均改善/%','胜/平/负电路数','旧耗时/s','新耗时/s'],
        [[r['objective'],LABELS[r['algorithm']],r['config'],f"{r['mean_best_gain_pct']:.2f}",f"{r['mean_mean_gain_pct']:.2f}",
         f"{r['improved_circuits']}/{r['tied_circuits']}/{r['worse_circuits']}",f"{r['baseline_mean_sec']:.2f}",f"{r['tuned_mean_sec']:.2f}"] for r in held])
    p('留出电路：bfly、dscg、ode、syn2。改善率以每个电路该算法上轮对应统计量为分母，再对四电路等权平均。负值表示退化，不进行截断。随机算法比较相同 seed0/1/2；MCTS 比较 seed0。该表才用于评价统一配置的迁移表现。')
    doc.add_page_break();h('4  选定配置与资源约束')
    table(['目标','算法','配置','训练比值','参数'],[[r['objective'],LABELS[r['algorithm']],r['selected_config'],f"{r['train_ratio']:.5f}",r['options']] for r in cfgrows])
    p('参数名与原生命令一致。baseline 为上一轮配置。候选 long20 使用长度20，并降低单步/回合预算；budget2 增加原长度下的搜索预算。候选明细、逐轮排名及锁定时刻见 manifest.json、promoted.json、selection.json 和 runner.log。')
    doc.add_page_break();h('5  复现、边界与结论使用')
    p('二进制 SHA-256：'+manifest['abc_sha256'])
    p('用户级 MemoryHigh='+str(memory['high_gib'])+' GiB，MemoryMax='+str(memory['max_gib'])+' GiB。'+notes[4][1])
    p(f"监测到的交换内存范围为 {memory['min_swap_mib']:.2f}–{memory['max_swap_mib']:.2f} MiB；这是整个用户 slice 的值，不仅包含本实验。")
    p('运行命令：python3 tools/synsearch/tune_vtr.py --previous results --out tuning_20261018 --max-workers 8。脚本按配置缓存已完成运行，重启前需确认无残留同目录任务。内存保护终止只作用于本脚本启动的进程组，重试记录保留。')
    p('本轮没有修改算法核心或动作集合。调参结果不能证明全局最优，也不能直接外推到 EPFL 等未测试数据集。算法的预算、序列长度与真实耗时不同，若用于论文，应同时呈现固定参数留出结果、成本信息与累计最好网络，避免只展示配置搜索后的最大收益。')
    p('原始数据保存在 '+notes[8][1]+'。每个新运行含命令、网络、原生JSON、资源日志、CEC日志；旧配置复用上轮结果，明确标记 origin=previous，没有伪造重复运行。')
    doc.save(out/'VTR调参_实验报告.docx')
    summary=dict(completed=completed,memory=memory,overview=overview,selected=selected,
                 holdout_aggregate=held,trial_rows=len(raw),new_successful=sum(r['origin']=='new' and r['status']=='ok' for r in trials))
    (out/'汇总.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False))
    print(json.dumps(summary,indent=2,ensure_ascii=False))

if __name__=='__main__':main()
